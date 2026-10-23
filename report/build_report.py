#!/usr/bin/env python3
"""Build the reviewed APA 7 student report from code evidence and local timings."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "report" / "Advanced_Shell_Deliverable_1_Report.docx"
TITLE = "Advanced Shell Simulation with Integrated OS Concepts"
SUBTITLE = "Deliverable 1 Basic Shell Implementation and Process Management"


def font(run, *, bold=False, italic=False, size=12):
    run.font.name = "Times New Roman"
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attribute), "Times New Roman")
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attribute), None)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold
    run.italic = italic
    return run


def paragraph(doc, text="", *, indent=True, bold=False, italic=False,
              centered=False, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.5 if indent else 0)
    pf.left_indent = Inches(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 2
    pf.widow_control = True
    pf.keep_with_next = keep
    font(p.add_run(text), bold=bold, italic=italic)
    return p


def heading(doc, text, level=1):
    return paragraph(doc, text, indent=False, bold=True, centered=level == 1,
                     keep=True, style=f"Heading {level}")


def page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)


def caption(doc, label, title):
    paragraph(doc, label, indent=False, bold=True, keep=True)
    paragraph(doc, title, indent=False, italic=True, keep=True)


def note(doc, text):
    p = paragraph(doc, indent=False)
    font(p.add_run("Note. "), italic=True)
    font(p.add_run(text))
    return p


def table(doc, label, title, headers, rows, widths):
    caption(doc, label, title)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    # APA 7 overrides the general document preset: horizontal rules only.
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        border.set(qn("w:val"), "single" if side in ("top", "bottom") else "nil")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    props.append(borders)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers] + rows):
        row = t.rows[0] if row_index == 0 else t.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, padding in (("top", 65), ("bottom", 65), ("left", 95), ("right", 95)):
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), str(padding))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            cell_props.append(margins)
            if row_index == 0:
                header_borders = OxmlElement("w:tcBorders")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "000000")
                header_borders.append(bottom)
                cell_props.append(header_borders)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or index == len(headers)-1 or (len(headers) == 4 and index > 0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            font(p.add_run(value), bold=row_index == 0)
    return t


def figure(doc, number, title, filename, description):
    caption(doc, f"Figure A{number}", title)
    p = paragraph(doc, indent=False)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(ROOT / "screenshots" / filename), width=Inches(6.5))
    picture._inline.docPr.set("descr", description)
    note(doc, description)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.5)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        properties = style._element.get_or_add_pPr()
        border = properties.find(qn("w:pBdr"))
        if border is not None:
            properties.remove(border)
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1
    r = p.add_run()
    font(r)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    p._p.append(field)
    doc.core_properties.author = "Shobhit Dhamania"
    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE


def main():
    metrics = json.loads((ROOT / "report" / "benchmark_results.json").read_text())
    doc = Document()
    configure(doc)

    # Student title page: same 12-point font, one blank line before the byline.
    p = paragraph(doc, TITLE + "\n" + SUBTITLE, indent=False,
                  centered=True, bold=True, style="Title")
    p.paragraph_format.space_before = Pt(72)
    paragraph(doc, indent=False)
    for text in ("Shobhit Dhamania", "University of the Cumberlands",
                 "MSCS-630-M20: Advanced Operating Systems", "Sean Perryman", "October 23, 2026"):
        paragraph(doc, text, indent=False, centered=True)
    page(doc)

    paragraph(doc, TITLE + "\n" + SUBTITLE, indent=False, centered=True,
              bold=True, keep=True, style="Heading 1")
    paragraph(doc,
        "Building this shell made process management easier for me to reason about because every command exposed an operating-system decision. A foreground program needs terminal access, a stopped program needs a state that survives the return of the prompt, and an exited child still needs its status collected. I implemented these behaviors in a C++17 shell and verified them through file operations, external programs, and interactive signal tests. The central result is a working Deliverable 1 implementation with 12 required utility built-ins and the three job-control commands jobs, fg, and bg.")
    paragraph(doc,
        "I used the POSIX specifications to examine the relationship between process creation, program replacement, and terminal control. The resulting shell owns its parser, built-ins, job table, launch ordering, and status handling. The host kernel still supplies actual processes, filesystem access, and CPU scheduling. This distinction matters: the project demonstrates shell-level control of OS mechanisms; it does not claim to replace the kernel or implement the scheduling, memory, synchronization, and security simulations reserved for subsequent deliverables.")

    heading(doc, "Design and Operating System Concepts")
    heading(doc, "Commands and Shell State", 2)
    paragraph(doc,
        "The command parser recognizes whitespace, single and double quotes, escaped characters, and a final unquoted ampersand. Ordinary backslashes inside double quotes remain literal. It rejects unmatched quotes, unfinished escapes, embedded NUL bytes, and unsupported operators before dispatch. CommandParser handles syntax, Shell coordinates execution, and Job stores an ID, PID, process-group ID, original command, status, and any saved terminal settings. Appendix B maps all 15 commands to their behavior and verification evidence.")
    paragraph(doc,
        "The cd command gave me a useful way to test process boundaries. A foreground cd must run in the parent so its directory change persists for the next prompt. A background cd runs in a child and leaves the parent's directory unchanged. Other built-ins can also run with an ampersand in a child; fg and bg reject that syntax because they must control the parent shell's jobs. In batch mode, background jobs read from /dev/null so they cannot consume the remaining shell commands.")

    heading(doc, "Process Creation and Terminal Ownership", 2)
    paragraph(doc,
        "POSIX specifies that fork creates a child with a separate PID, whereas exec replaces the current process image (Institute of Electrical and Electronics Engineers [IEEE] & The Open Group, 2024a, 2024b). That difference clarified the launch sequence for me: fork creates the child, and execvp turns that child into the requested program. An external command is therefore executed directly through an argument vector and PATH lookup. The shell does not pass every line to system() or hide its process-management logic inside another shell.")
    paragraph(doc,
        "The parent and child establish a process group before program execution. A launch pipe holds the child until the parent has registered a background job or assigned the terminal to a foreground job. A second pipe, whose child endpoint closes on exec, acknowledges the launch transition before the shell accepts another command. These ordering steps address concrete races: a foreground reader could otherwise receive SIGTTIN before owning the terminal, or an immediate exit could overlap a child's program replacement.")
    paragraph(doc,
        "The foreground group is assigned with tcsetpgrp, an interface explicitly defined for controlling-terminal ownership (IEEE & The Open Group, 2018a). The shell then waits for status changes. After a stop or exit, it restores its own foreground group and terminal settings. The tests verify this with an actual terminal reader, so success requires both the right process group and usable input, rather than only a job-table label.")

    heading(doc, "Job States and Signal Safety", 2)
    paragraph(doc,
        "A job is recorded as Running, Stopped, or Done. These are shell-level summaries: Running means the job has not stopped or completed, and may include a process that the kernel currently puts to sleep. A Ctrl-Z stop retains the job and its terminal settings. bg sends SIGCONT to its group; fg restores terminal ownership and saved settings before continuing it. This ordering follows POSIX job-control requirements (IEEE & The Open Group, 2024c). Figure A3 demonstrates the complete sequence, including Ctrl-C termination after fg.")
    paragraph(doc,
        "The SIGCHLD handler writes one byte to a nonblocking self-pipe and preserves errno. The normal input loop polls that pipe alongside standard input, then calls waitpid to update the job table. This avoids C++ stream output and container mutation inside a signal handler, consistent with POSIX signal-safety guidance (IEEE & The Open Group, 2024d). A full pipe already provides a wakeup; the code therefore drains available child statuses rather than assuming one notification equals one child.")
    paragraph(doc,
        "waitpid reports completion, stops, and continuations, with WNOHANG allowing status checks without blocking (IEEE & The Open Group, 2018b). The foreground wait path also collects other children's statuses so a background completion is not postponed until a long foreground command ends. At the prompt, the self-pipe wakes status collection even when the user types nothing. On shell exit, active groups receive SIGHUP and stopped groups also receive SIGCONT so they can respond to the hangup.")

    heading(doc, "Verification and Performance")
    heading(doc, "Functional Evidence", 2)
    paragraph(doc,
        "The strict build enables C++17 and treats compiler warnings as errors. Parser assertions check argument grouping and invalid syntax. Twenty integration tests check filesystem operations, exit codes, signals, child cleanup, and job control. A pseudo-terminal supplies a real controlling terminal for the interactive tests. The same 20-test suite passes with AddressSanitizer and UndefinedBehaviorSanitizer enabled, and Clang static analysis reports no findings. These checks cover the exercised paths; they are not proof that every possible input is correct.")
    paragraph(doc,
        "The expanded tests include an idle background completion, 12 successive foreground terminal readers, a reader stopped by SIGTTIN and resumed by fg, and a nested shell initially launched in the background. A dedicated test changes a job's echo setting before stopping it, then verifies that fg restores that setting. These cases distinguish execution correctness from superficial output correctness. Both PID tracking and terminal behavior must agree for a test to pass.")
    paragraph(doc,
        "Appendix A provides five terminal-session captures with their original text transcripts in the submission. Figure A1 covers file and directory commands, including empty-file output. Figure A2 independently checks a background PID and PGID with ps and terminates that process through kill. Figure A3 demonstrates fg and bg. Figure A4 shows invalid commands and inputs. Figure A5 adds clear, a background built-in, and exit. The PNGs render the recorded PTY output; they are reproducible session evidence rather than photographs of a native Terminal window.")

    heading(doc, "Performance Findings and Limits", 2)
    paragraph(doc,
        "I measured complete batch runs with the optimized executable on macOS 15.7.8, Apple Silicon, and Apple clang 17.0.0. Each workload ran nine times with output discarded. Table 1 reports the median and observed range, including shell startup. The benchmark script and all raw timings are included, allowing the comparison to be repeated instead of relying on an unrecorded timing claim.")
    table(doc, "Table 1", "Elapsed Times for Repeated Shell Workloads",
          ["Workload", "Count", "Median ms", "Range ms"],
          [["Startup and exit", "0", f"{metrics['baseline']['median_ms']:.2f}",
            f"{metrics['baseline']['minimum_ms']:.2f}–{metrics['baseline']['maximum_ms']:.2f}"],
           ["Built-in cd .", "1,000", f"{metrics['builtin']['median_ms']:.2f}",
            f"{metrics['builtin']['minimum_ms']:.2f}–{metrics['builtin']['maximum_ms']:.2f}"],
           ["External /usr/bin/true", "200", f"{metrics['external']['median_ms']:.2f}",
            f"{metrics['external']['minimum_ms']:.2f}–{metrics['external']['maximum_ms']:.2f}"]],
          [2.8, 0.8, 1.2, 1.7])
    note(doc, "Each value summarizes nine runs. ms = milliseconds. Ranges show the minimum and maximum; they are not confidence intervals.")
    builtin_per = (metrics['builtin']['median_ms'] - metrics['baseline']['median_ms']) / 1000
    external_per = (metrics['external']['median_ms'] - metrics['baseline']['median_ms']) / 200
    paragraph(doc,
        f"Subtracting the baseline median gives rough per-command estimates of {builtin_per * 1000:.1f} microseconds for cd and {external_per:.2f} milliseconds for the external program. I interpret the difference as the cost of taking the process-launch path, including pipes, fork, exec, waiting, and kernel scheduling. The experiment does not isolate any one system call. Subtracting medians is also an approximation, and the short built-in workload is especially sensitive to startup variation. These measurements support the design tradeoff without establishing a universal speed ratio or a Linux performance result.")

    heading(doc, "Error Handling and Design Revisions")
    paragraph(doc,
        "Errors are reported with the command, target, and reason where available. Missing arguments and invalid PIDs fail before filesystem or signal operations. A failed exec exits the child with 127 for a missing executable or 126 for another launch error. Parsing errors return status 2; ordinary built-in failures return 1. An explicit exit status is honored, while exit without a status and end-of-input preserve the last command result. An error therefore remains observable to a caller as well as to someone reading the screen.")
    paragraph(doc,
        "Reviewing boundary cases exposed an output bug that ordinary demonstrations had missed. Streaming an empty file buffer into std::cout could mark the output stream as failed and suppress subsequent prompts or echo output. The revised cat copies bytes explicitly, accepts an empty file without changing output state, and rejects directory input. Its regression test requires successful output after both an empty file and a directory error. That turned a seemingly harmless empty input into a useful lesson about state carried across commands.")
    paragraph(doc,
        "Filesystem tests also exposed a distinction between a directory and a link to one. rmdir now uses the OS directory-removal operation, which rejects a symlink, while rm checks the link itself and can unlink it without deleting its target. touch updates existing timestamps and creates missing files without truncating data. The tests preserve sample file contents, reject removal of nonempty directories, and verify that directory targets survive symlink removal.")
    paragraph(doc,
        "The job-control review found two more timing problems. Ignoring SIGTTIN before checking foreground access could prevent a nested shell from stopping correctly during startup. Collecting child status only between input lines could leave an exited background process waiting indefinitely for another keystroke. The revised initialization order and self-pipe address those causes directly. Further regression tests caught a batch background reader consuming later commands and a quoted path losing a backslash. Isolating batch input and preserving literal quoted characters corrected those separate failures.")

    heading(doc, "Reflection and Future Improvements")
    paragraph(doc,
        "The moment that made the project click for me was watching one sleep process move from foreground execution to Stopped, resume under bg, and return through fg before Ctrl-C terminated it. The PID persisted across those transitions. That made a textbook process-state discussion tangible: stopping changes a live process's state, while terminating ends its execution. The shell's job table became a record of a lifecycle I could actually observe and control.")
    paragraph(doc,
        "Reading the POSIX job-control section helped explain why fg needs more than waitpid. Its behavior depends on the terminal's foreground group and on restoring a stopped program's terminal settings (IEEE & The Open Group, 2024c). The terminal-mode test was especially satisfying because it verified something a normal sleep demonstration cannot show. A prompt and a PID looked sufficient at first, but input ownership and saved state explained why a shell can safely support interactive applications.")
    paragraph(doc,
        "The error cases changed what I consider a convincing demonstration. A successful cat of a nonempty file did not establish that cat was safe, and a jobs listing did not establish that exited children were collected while the prompt was idle. I found the project more interesting when the tests began challenging those assumptions. Each failure connected a visible symptom to a specific mechanism, which made the OS concepts feel less abstract and made the final design easier to defend.")
    paragraph(doc,
        "The next engineering priority is to run this same suite on Linux; the source uses portable POSIX APIs, but the completed execution evidence is macOS-only. Pipelines would require a job to own multiple PIDs and aggregate their states, while preserving the existing process-group boundary. Redirection, full shell expansion, and editable history are further extensions. The later project deliverables can introduce simulated scheduling, memory management, synchronization, and security policies, with this shell serving as their command interface. Current commands operate on the real filesystem under the user's existing permissions; this deliverable is not a security sandbox.")

    heading(doc, "Conclusion")
    paragraph(doc,
        "The completed shell meets Deliverable 1 through implemented built-ins, direct external execution, foreground and background jobs, status tracking, and error feedback. The review strengthened the result by fixing failures involving empty files, symbolic links, idle status collection, and terminal ownership. My main learning outcome is a more precise understanding of how a shell coordinates kernel mechanisms: the process API creates opportunities for concurrency, and the shell must impose the ordering and state transitions that make those opportunities usable.")

    page(doc)
    heading(doc, "References")
    references = [
        ("2018a", "tcsetpgrp", "The Open Group base specifications issue 7, 2018 edition",
         "https://pubs.opengroup.org/onlinepubs/9699919799/functions/tcsetpgrp.html"),
        ("2018b", "wait", "The Open Group base specifications issue 7, 2018 edition",
         "https://pubs.opengroup.org/onlinepubs/9699919799/functions/wait.html"),
        ("2024a", "exec", "The Open Group base specifications issue 8",
         "https://pubs.opengroup.org/onlinepubs/9799919799/functions/exec.html"),
        ("2024b", "fork", "The Open Group base specifications issue 8",
         "https://pubs.opengroup.org/onlinepubs/9799919799/functions/fork.html"),
        ("2024c", "Shell command language", "The Open Group base specifications issue 8",
         "https://pubs.opengroup.org/onlinepubs/9799919799/utilities/V3_chap02.html"),
        ("2024d", "sigaction", "The Open Group base specifications issue 8",
         "https://pubs.opengroup.org/onlinepubs/9799919799/functions/sigaction.html"),
    ]
    for year, title, collection, url in references:
        p = paragraph(doc, indent=False)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        font(p.add_run(f"Institute of Electrical and Electronics Engineers & The Open Group. ({year}). {title}. In "))
        font(p.add_run(collection), italic=True)
        font(p.add_run(f". {url}"))

    page(doc)
    heading(doc, "Appendix A")
    heading(doc, "Demonstration Evidence")
    figure(doc, 1, "File and Directory Built-ins", "01_builtins.png",
           "Recorded PTY output demonstrates pwd, echo, mkdir, cd, touch, ls, cat, rm, and rmdir. The sample text is a prepared fixture. Output continues after cat reads the empty notes.txt file.")
    page(doc)
    figure(doc, 2, "Background Process Creation and Termination", "02_background_and_kill.png",
           "sleep runs asynchronously. jobs displays its shell state and group ID; ps independently confirms the PID and PGID. The built-in kill sends SIGTERM, and the completed child is reported as Done. PID = process ID; PGID = process-group ID.")
    page(doc)
    figure(doc, 3, "Stop Resume and Foreground Terminal Control", "03_fg_bg_signals.png",
           "Ctrl-Z stops the same process, bg resumes its group, and fg returns it to the foreground. Ctrl-C then terminates it while the shell remains available for another command.")
    page(doc)
    figure(doc, 4, "Invalid Commands and Input Errors", "04_error_handling.png",
           "The session exercises missing commands and files, invalid directories, a nonempty directory, an invalid PID, unmatched quotes, an unsupported pipe, and a missing argument. The shell remains responsive.")
    page(doc)
    figure(doc, 5, "Clear Background Built-in and Shell Exit", "05_clear_and_exit.png",
           "clear emits its ANSI control sequence, which is omitted from the text rendering. The integration test checks its exact bytes. echo executes in a background child, and exit 0 ends the shell. ANSI = American National Standards Institute.")

    page(doc)
    heading(doc, "Appendix B")
    heading(doc, "Command Coverage and Submission Contents")
    paragraph(doc, "Table B1 maps every required command to implementation behavior and direct verification. The source, Makefile, tests, README, benchmark data, screenshots, and transcripts accompany this report. Build with make, run ./build/os-shell, and verify with make test. GitHub publication is deferred.")
    table(doc, "Table B1", "Required Command Coverage",
          ["Command", "Implemented behavior", "Evidence"],
          [["cd [directory]", "Changes parent cwd; no argument uses HOME.", "A1; integration tests"],
           ["pwd", "Prints current directory.", "A1"],
           ["exit [status]", "Exits shell; validates optional status.", "A5; exit tests"],
           ["echo [text]", "Prints parsed arguments.", "A1; A5"],
           ["clear", "Emits ANSI clear and cursor-home codes.", "A5; byte test"],
           ["ls [path]", "Lists sorted nonhidden entries.", "A1"],
           ["cat file...", "Copies bytes; handles empty files and read errors.", "A1; regression tests"],
           ["mkdir directory...", "Creates directories; reports failures.", "A1; A4"],
           ["rmdir directory...", "Removes empty directories; rejects symlinks.", "A1; A4; link tests"],
           ["rm file...", "Removes files or links; refuses directories.", "A1; A4; link tests"],
           ["touch file...", "Creates missing files; updates timestamps.", "A1; timestamp test"],
           ["kill pid", "Sends SIGTERM to a validated positive PID.", "A2; A4"],
           ["jobs", "Lists tracked running, stopped, and done jobs.", "A2; A3"],
           ["fg [job_id]", "Restores terminal ownership and resumes/waits.", "A3; terminal tests"],
           ["bg [job_id]", "Resumes the stopped group asynchronously.", "A3"]],
          [1.3, 3.65, 1.55])
    note(doc, "A1–A5 refer to Figures A1–A5. cwd = current working directory. Both numeric job IDs and percent-prefixed IDs are accepted. With no ID, fg and bg use the most recent active job.")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
